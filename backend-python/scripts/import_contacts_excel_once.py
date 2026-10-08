"""Import a WeCom address-book Excel export into HRM.

The script is intentionally a one-time, explicit import tool. It defaults to a
dry run and only writes when ``--apply`` is passed.

Expected workbook shape:
  - A header row containing at least: 姓名, 账号, 部门
  - WeCom-style department paths separated by "/" and multiple memberships by ";"
"""

from __future__ import annotations

import argparse
import hashlib
import os
import secrets
import sys
import uuid
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Iterable

import bcrypt
from dotenv import load_dotenv
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

load_dotenv(ROOT.parent / '.env')
os.environ.setdefault('DISABLE_STUDY_REMINDER_SCHEDULER', '1')

from app import create_app  # noqa: E402
from models import Department, Employee, EmployeePortalAccount, User, db  # noqa: E402
from services.employee_portal_accounts import ensure_employee_portal_account  # noqa: E402


NAMESPACE = uuid.uuid5(uuid.NAMESPACE_URL, 'hengpintong:wecom-excel-contacts')
SOURCE_LABEL = '企业微信通讯录 Excel 一次性导入'
IMPORT_MARKER = '[wecom-excel-import]'
REQUIRED_HEADERS = {'姓名', '账号', '部门'}
OPTIONAL_HEADERS = {
    '别名',
    '职务',
    '性别',
    '手机',
    '座机',
    '邮箱',
    '地址',
    '企业邮箱',
    '激活状态',
    '禁用状态',
}


def _stable_id(kind: str, source_id: str) -> str:
    return str(uuid.uuid5(NAMESPACE, f'{kind}:{source_id}'))


def _hash(value: str, length: int = 12) -> str:
    return hashlib.sha1(value.encode('utf-8')).hexdigest()[:length]


def _clean(value: object) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()


def _limit(value: str | None, max_length: int) -> str | None:
    if not value:
        return None
    return value[:max_length]


def _safe_account(account: str) -> str:
    safe = ''.join(ch.lower() if ch.isalnum() or ch in '._-' else '-' for ch in account)
    safe = safe.strip('.-_')
    return safe or f'user-{_hash(account)}'


def _employee_no(account: str) -> str:
    raw = f'WX-{account}'
    if len(raw) <= 50:
        return raw
    return f'WX-{_hash(account, 16)}'


def _department_code(path: tuple[str, ...]) -> str:
    return f'excel:{_hash("/".join(path), 16)}'


def _gender(value: str) -> str | None:
    normalized = value.strip()
    if normalized in {'男', '女'}:
        return normalized
    if normalized in {'其他', '其它'}:
        return '其他'
    return None


def _status(activation_state: str, disabled_state: str) -> str:
    disabled = disabled_state.strip().lower()
    if disabled in {'1', 'true', 'yes', 'y', '是', '已禁用', '禁用'}:
        return 'inactive'
    activation = activation_state.strip()
    if activation and activation not in {'已激活', '激活', '正常'}:
        return 'inactive'
    return 'active'


def _split_department_paths(raw: str) -> list[tuple[str, ...]]:
    paths: list[tuple[str, ...]] = []
    for item in raw.split(';'):
        parts = tuple(part.strip() for part in item.replace('\\', '/').split('/') if part.strip())
        if parts:
            paths.append(parts)
    return paths


def _find_header_row(ws, forced_row: int | None) -> tuple[int, dict[str, int]]:
    max_scan_row = min(ws.max_row or 40, 40)
    rows: Iterable[int] = [forced_row] if forced_row else range(1, max_scan_row + 1)
    for row_number in rows:
        values = [_clean(cell.value) for cell in ws[row_number]]
        headers = {value: index for index, value in enumerate(values) if value}
        if REQUIRED_HEADERS.issubset(headers.keys()):
            return row_number, headers
    raise RuntimeError('未找到企业微信通讯录表头行，至少需要包含：姓名、账号、部门')


def _load_contacts(path: Path, sheet_name: str | None, header_row: int | None) -> tuple[list[dict], dict]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    ws = workbook[sheet_name] if sheet_name else workbook.active
    row_number, headers = _find_header_row(ws, header_row)
    known_headers = REQUIRED_HEADERS | OPTIONAL_HEADERS
    available_headers = sorted(set(headers) & known_headers)

    contacts: list[dict] = []
    skipped_rows = 0
    for excel_row, row in enumerate(ws.iter_rows(min_row=row_number + 1, values_only=True), start=row_number + 1):
        if not any(_clean(value) for value in row):
            continue
        data = {header: _clean(row[index]) if index < len(row) else '' for header, index in headers.items()}
        name = data.get('姓名', '')
        account = data.get('账号', '')
        department_raw = data.get('部门', '')
        if not (name and account and department_raw):
            skipped_rows += 1
            continue
        paths = _split_department_paths(department_raw)
        if not paths:
            skipped_rows += 1
            continue
        contacts.append(
            {
                'row': excel_row,
                'name': name,
                'account': account,
                'alias': data.get('别名', ''),
                'job_title': data.get('职务', ''),
                'department_raw': department_raw,
                'department_paths': paths,
                'gender': data.get('性别', ''),
                'mobile': data.get('手机', ''),
                'telephone': data.get('座机', ''),
                'email': data.get('邮箱', ''),
                'biz_email': data.get('企业邮箱', ''),
                'address': data.get('地址', ''),
                'activation_state': data.get('激活状态', ''),
                'disabled_state': data.get('禁用状态', ''),
            }
        )

    return contacts, {
        'sheet_name': ws.title,
        'header_row': row_number,
        'available_headers': available_headers,
        'skipped_rows': skipped_rows,
    }


def _email_for_contact(contact: dict, synthetic_domain: str, duplicate_emails: set[str]) -> tuple[str, bool]:
    raw_email = (contact.get('email') or contact.get('biz_email') or '').strip().lower()
    if raw_email and '@' in raw_email and raw_email not in duplicate_emails:
        return raw_email, False
    account = _safe_account(contact['account'])
    return f'{account}@{synthetic_domain}', True


def _password_hash(default_password: str | None) -> str:
    password = default_password or secrets.token_urlsafe(48)
    return bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')


def _replace_import_note(existing: str | None, note: str) -> str:
    lines = [line for line in (existing or '').splitlines() if not line.startswith(IMPORT_MARKER)]
    lines.append(f'{IMPORT_MARKER} {note}')
    return '\n'.join(line for line in lines if line).strip()


def _department_by_path(path: tuple[str, ...], parent_id: str | None) -> Department | None:
    dept_id = _stable_id('department', '/'.join(path))
    return (
        db.session.get(Department, dept_id)
        or Department.query.filter_by(code=_department_code(path)).first()
        or Department.query.filter_by(name=path[-1], parent_id=parent_id).first()
    )


def _upsert_departments(contacts: list[dict], stats: Counter) -> dict[tuple[str, ...], str]:
    all_paths: set[tuple[str, ...]] = set()
    for contact in contacts:
        for path in contact['department_paths']:
            for index in range(1, len(path) + 1):
                all_paths.add(path[:index])

    dept_id_by_path: dict[tuple[str, ...], str] = {}
    for path in sorted(all_paths, key=lambda item: (len(item), item)):
        parent_id = dept_id_by_path.get(path[:-1]) if len(path) > 1 else None
        dept = _department_by_path(path, parent_id)
        if dept:
            stats['departments_updated'] += 1
        else:
            dept = Department(
                id=_stable_id('department', '/'.join(path)),
                code=_department_code(path),
                description=f'{SOURCE_LABEL}：{"/".join(path)}',
            )
            db.session.add(dept)
            stats['departments_created'] += 1
        dept.name = _limit(path[-1], 100)
        if not dept.code:
            dept.code = _department_code(path)
        dept.parent_id = parent_id
        dept.status = 'active'
        dept.updated_at = datetime.utcnow()
        dept_id_by_path[path] = dept.id
    db.session.flush()
    return dept_id_by_path


def _employee_by_contact(contact: dict, email: str) -> Employee | None:
    employee_id = _stable_id('employee', contact['account'])
    employee_no = _employee_no(contact['account'])
    return (
        db.session.get(Employee, employee_id)
        or Employee.query.filter_by(employee_no=employee_no).first()
        or Employee.query.filter_by(email=email).first()
    )


def _user_by_contact(contact: dict, email: str) -> User | None:
    return db.session.get(User, _stable_id('user', contact['account'])) or User.query.filter_by(
        email=email
    ).first()


def import_contacts(
    path: Path,
    *,
    sheet_name: str | None,
    header_row: int | None,
    apply: bool,
    create_users: bool,
    default_role: str,
    default_password: str | None,
    synthetic_domain: str,
) -> dict[str, int | str | bool]:
    contacts, metadata = _load_contacts(path, sheet_name, header_row)
    stats: Counter = Counter()
    stats['excel_rows'] = len(contacts)
    stats['skipped_rows'] = metadata['skipped_rows']
    stats['multi_department_contacts'] = sum(1 for contact in contacts if len(contact['department_paths']) > 1)
    stats['unique_department_paths'] = len({path for contact in contacts for path in contact['department_paths']})

    email_counts = Counter(
        (contact.get('email') or contact.get('biz_email') or '').strip().lower()
        for contact in contacts
        if (contact.get('email') or contact.get('biz_email') or '').strip()
    )
    duplicate_emails = {email for email, count in email_counts.items() if count > 1}
    stats['duplicate_emails_synthesized'] = len(duplicate_emails)

    dept_id_by_path = _upsert_departments(contacts, stats)

    for contact in contacts:
        primary_path = contact['department_paths'][0]
        local_dept_id = dept_id_by_path.get(primary_path)
        email, synthesized_email = _email_for_contact(contact, synthetic_domain, duplicate_emails)
        if synthesized_email:
            stats['synthetic_emails_used'] += 1

        status = _status(contact.get('activation_state', ''), contact.get('disabled_state', ''))
        if status != 'active':
            stats['inactive_contacts'] += 1

        employee = _employee_by_contact(contact, email)
        if employee:
            stats['employees_updated'] += 1
        else:
            employee = Employee(id=_stable_id('employee', contact['account']), employee_no=_employee_no(contact['account']))
            db.session.add(employee)
            stats['employees_created'] += 1

        employee.name = _limit(contact['name'], 100)
        employee.employee_no = employee.employee_no or _employee_no(contact['account'])
        employee.email = _limit(email, 255)
        employee.phone = _limit(contact.get('mobile') or contact.get('telephone'), 30)
        employee.gender = _gender(contact.get('gender', '')) or employee.gender
        employee.status = status
        employee.department_id = local_dept_id
        employee.job_title = _limit(contact.get('job_title'), 120)
        employee.address = _limit(contact.get('address'), 255)
        employee.remark = _replace_import_note(
            employee.remark,
            f'账号={contact["account"]}; 部门={contact["department_raw"]}; 来源={SOURCE_LABEL}',
        )
        employee.updated_at = datetime.utcnow()

        had_portal_account = EmployeePortalAccount.query.filter_by(employee_id=employee.id).first() is not None
        ensure_employee_portal_account(
            employee,
            initial_password=default_password,
            reset_password=not had_portal_account,
            require_password_change=True,
        )
        stats['portal_accounts_updated' if had_portal_account else 'portal_accounts_created'] += 1

        if create_users:
            account = _user_by_contact(contact, email)
            if account:
                stats['users_updated'] += 1
            else:
                account = User(
                    id=_stable_id('user', contact['account']),
                    email=email,
                    password_hash=_password_hash(default_password),
                    must_change_password=True,
                    role=default_role,
                )
                db.session.add(account)
                stats['users_created'] += 1
            account.name = _limit(contact['name'], 100)
            account.phone = _limit(contact.get('mobile') or contact.get('telephone'), 30)
            account.status = status
            account.department_id = local_dept_id
            account.employee_id = employee.id
            account.updated_at = datetime.utcnow()

    if apply:
        db.session.commit()
    else:
        db.session.rollback()

    result: dict[str, int | str | bool] = {key: int(value) for key, value in sorted(stats.items())}
    result['mode'] = 'APPLY' if apply else 'DRY-RUN'
    result['sheet_name'] = metadata['sheet_name']
    result['header_row'] = metadata['header_row']
    result['create_users'] = create_users
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description='Import WeCom contacts Excel into HRM')
    parser.add_argument('xlsx', type=Path, help='WeCom address-book .xlsx file')
    parser.add_argument('--sheet', dest='sheet_name', help='worksheet name; defaults to the active sheet')
    parser.add_argument('--header-row', type=int, help='header row number; auto-detected when omitted')
    parser.add_argument('--apply', action='store_true', help='write to database; default is dry-run')
    parser.add_argument('--no-create-users', action='store_true', help='only import departments and employees')
    parser.add_argument('--default-role', default=os.getenv('WECOM_DEFAULT_ROLE', 'employee'))
    parser.add_argument('--default-password', default=os.getenv('WECOM_DEFAULT_PASSWORD'))
    parser.add_argument(
        '--synthetic-email-domain',
        default=os.getenv('WECOM_SYNTHETIC_EMAIL_DOMAIN', 'wecom.local'),
    )
    args = parser.parse_args()

    if not args.xlsx.exists():
        raise RuntimeError(f'Excel file does not exist: {args.xlsx}')

    app = create_app()
    with app.app_context():
        stats = import_contacts(
            args.xlsx,
            sheet_name=args.sheet_name,
            header_row=args.header_row,
            apply=args.apply,
            create_users=not args.no_create_users,
            default_role=(args.default_role or 'employee').strip() or 'employee',
            default_password=args.default_password,
            synthetic_domain=(args.synthetic_email_domain or 'wecom.local').strip() or 'wecom.local',
        )

    for key in sorted(stats):
        print(f'{key}={stats[key]}')
    if not args.apply:
        print('dry_run=true; pass --apply to write changes')
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
